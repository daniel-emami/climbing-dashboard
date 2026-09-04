from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from ClimbingDashboard.Config.constants import BOULDER_EXPORT_COLUMNS
from ClimbingDashboard.Models.boulder_record import BoulderRecord


class ExcelBoulderExporter:
    """Build Excel workbooks from boulder records."""

    def build_workbook(self, boulders: list[BoulderRecord]) -> BytesIO:
        """Return an in-memory workbook for the supplied boulder records."""

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Boulders"
        for index, header in enumerate(BOULDER_EXPORT_COLUMNS, start=1):
            worksheet.cell(1, index, header)

        for row_index, boulder in enumerate(boulders, start=2):
            worksheet.cell(row_index, 1, boulder.name)
            worksheet.cell(row_index, 2, boulder.grade_27crags)
            worksheet.cell(row_index, 3, boulder.guide_grade)
            worksheet.cell(row_index, 4, boulder.own_grade)
            worksheet.cell(row_index, 5, boulder.area)
            worksheet.cell(row_index, 6, 1 if boulder.flash else 0)
            date_cell = worksheet.cell(row_index, 7, boulder.climbed_on)
            date_cell.number_format = "yyyy-mm-dd"
            worksheet.cell(row_index, 8, boulder.climber)
            worksheet.cell(row_index, 9, boulder.rating)
            worksheet.cell(row_index, 10, boulder.visibility)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        return stream
