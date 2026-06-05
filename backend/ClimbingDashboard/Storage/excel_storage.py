from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ClimbingDashboard.Exceptions.excel_storage_error import ExcelStorageError
from ClimbingDashboard.Models.boulder_record import BoulderRecord
from ClimbingDashboard.Storage.base_storage import BaseStorage
from ClimbingDashboard.Utilities.date_utils import parse_excel_date, to_excel_date

logger = logging.getLogger(__name__)

SOURCE_COLUMNS = {
    "A": "Navn",
    "B": "27Crags grade",
    "C": "Guide grade",
    "D": "My grade",
    "E": "Område",
    "F": "Flash",
    "G": "Dato",
}


class ExcelStorage(BaseStorage):
    """Excel storage implementation using openpyxl."""

    def __init__(self, workbook_path: str | Path) -> None:
        """Create storage bound to one workbook path."""

        self.workbook_path = Path(workbook_path)

    def read_boulders(self) -> list[BoulderRecord]:
        """Read boulder rows from the first worksheet."""

        worksheet = self._load_worksheet()
        self._validate_headers(worksheet)
        records: list[BoulderRecord] = []
        for row in worksheet.iter_rows(min_row=2, max_col=7, values_only=True):
            if not any(row):
                continue
            records.append(self._row_to_record(row))
        return records

    def append_boulder(self, record: BoulderRecord) -> BoulderRecord:
        """Append one boulder record to the workbook and persist it."""

        self.append_boulders([record])
        return record

    def append_boulders(self, records: list[BoulderRecord]) -> list[BoulderRecord]:
        """Append boulder records to the workbook and persist them."""

        if not self.workbook_path.exists():
            raise ExcelStorageError(f"Workbook does not exist: {self.workbook_path}")
        try:
            workbook = load_workbook(self.workbook_path)
            worksheet = workbook.worksheets[0]
            self._validate_headers(worksheet)
            next_row = self._next_source_row(worksheet)
            existing_keys = self._existing_boulder_keys(worksheet)
            appended_records: list[BoulderRecord] = []
            for record in records:
                record_key = self._boulder_key(record.name, record.area)
                if record_key in existing_keys:
                    logger.info("Skipping duplicate boulder %s in %s", record.name, record.area)
                    continue
                self._write_record_row(worksheet, next_row, record)
                existing_keys.add(record_key)
                appended_records.append(record)
                next_row += 1
            workbook.save(self.workbook_path)
            logger.info("Appended %s boulders to %s", len(appended_records), self.workbook_path)
        except ExcelStorageError:
            raise
        except Exception as exc:
            raise ExcelStorageError(f"Failed to append boulders: {exc}") from exc
        return appended_records

    def update_boulder(
        self,
        original_name: str,
        original_area: str,
        record: BoulderRecord,
    ) -> BoulderRecord:
        """Update one boulder matched by its original name and area."""

        if not self.workbook_path.exists():
            raise ExcelStorageError(f"Workbook does not exist: {self.workbook_path}")
        try:
            workbook = load_workbook(self.workbook_path)
            worksheet = workbook.worksheets[0]
            self._validate_headers(worksheet)
            target_row = self._find_boulder_row(worksheet, original_name, original_area)
            if target_row is None:
                raise ExcelStorageError(
                    f"Boulder does not exist: {original_name} in {original_area}"
                )

            updated_key = self._boulder_key(record.name, record.area)
            existing_keys = self._existing_boulder_keys(worksheet, excluded_row=target_row)
            if updated_key in existing_keys:
                raise ExcelStorageError(f"Boulder already exists: {record.name} in {record.area}")

            self._write_record_row(worksheet, target_row, record)
            workbook.save(self.workbook_path)
            logger.info("Updated boulder %s in %s", record.name, record.area)
        except ExcelStorageError:
            raise
        except Exception as exc:
            raise ExcelStorageError(f"Failed to update boulder: {exc}") from exc
        return record

    def delete_boulder(self, name: str, area: str) -> None:
        """Delete one boulder matched by name and area."""

        if not self.workbook_path.exists():
            raise ExcelStorageError(f"Workbook does not exist: {self.workbook_path}")
        try:
            workbook = load_workbook(self.workbook_path)
            worksheet = workbook.worksheets[0]
            self._validate_headers(worksheet)
            target_row = self._find_boulder_row(worksheet, name, area)
            if target_row is None:
                raise ExcelStorageError(f"Boulder does not exist: {name} in {area}")
            worksheet.delete_rows(target_row, 1)
            workbook.save(self.workbook_path)
            logger.info("Deleted boulder %s in %s", name, area)
        except ExcelStorageError:
            raise
        except Exception as exc:
            raise ExcelStorageError(f"Failed to delete boulder: {exc}") from exc

    def _load_worksheet(self) -> Worksheet:
        if not self.workbook_path.exists():
            raise ExcelStorageError(f"Workbook does not exist: {self.workbook_path}")
        try:
            workbook = load_workbook(self.workbook_path, data_only=True)
            return workbook.worksheets[0]
        except Exception as exc:
            raise ExcelStorageError(f"Failed to read workbook: {exc}") from exc

    def _validate_headers(self, worksheet: Worksheet) -> None:
        for column, expected_header in SOURCE_COLUMNS.items():
            observed_header = worksheet[f"{column}1"].value
            if observed_header != expected_header:
                raise ExcelStorageError(
                    f"Expected header {expected_header!r} in {column}1, found {observed_header!r}"
                )

    def _row_to_record(self, row: tuple[Any, ...]) -> BoulderRecord:
        try:
            return BoulderRecord(
                name=self._text(row[0]),
                grade_27crags=self._text(row[1]),
                guide_grade=self._text(row[2]),
                my_grade=self._text(row[3]),
                area=self._text(row[4]),
                flash=self._bool(row[5]),
                climbed_on=parse_excel_date(row[6]),
            )
        except ValueError as exc:
            raise ExcelStorageError(str(exc)) from exc

    def _next_source_row(self, worksheet: Worksheet) -> int:
        row_number = 2
        while worksheet.cell(row_number, 1).value not in (None, ""):
            row_number += 1
        return row_number

    def _write_record_row(
        self,
        worksheet: Worksheet,
        row_number: int,
        record: BoulderRecord,
    ) -> None:
        worksheet.cell(row_number, 1, record.name)
        worksheet.cell(row_number, 2, record.grade_27crags)
        worksheet.cell(row_number, 3, record.guide_grade)
        worksheet.cell(row_number, 4, record.my_grade)
        worksheet.cell(row_number, 5, record.area)
        worksheet.cell(row_number, 6, 1 if record.flash else 0)
        date_cell = worksheet.cell(row_number, 7, to_excel_date(record.climbed_on))
        date_cell.number_format = "yyyy-mm-dd"

    def _existing_boulder_keys(
        self,
        worksheet: Worksheet,
        excluded_row: int | None = None,
    ) -> set[tuple[str, str]]:
        keys: set[tuple[str, str]] = set()
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, max_col=5, values_only=True),
            start=2,
        ):
            if row_number == excluded_row:
                continue
            name = self._text(row[0])
            area = self._text(row[4])
            if name and area:
                keys.add(self._boulder_key(name, area))
        return keys

    def _find_boulder_row(
        self,
        worksheet: Worksheet,
        name: str,
        area: str,
    ) -> int | None:
        target_key = self._boulder_key(name, area)
        for row_number in range(2, worksheet.max_row + 1):
            row_name = self._text(worksheet.cell(row_number, 1).value)
            row_area = self._text(worksheet.cell(row_number, 5).value)
            if row_name and row_area and self._boulder_key(row_name, row_area) == target_key:
                return row_number
        return None

    def _boulder_key(self, name: str, area: str) -> tuple[str, str]:
        return (name.strip().casefold(), area.strip().casefold())

    def _text(self, value: object) -> str:
        return "" if value is None else str(value).strip()

    def _bool(self, value: object) -> bool:
        if isinstance(value, bool):
            return value
        if isinstance(value, int | float):
            return bool(value)
        if isinstance(value, str):
            return value.strip().lower() in {"1", "true", "yes", "y", "flash"}
        return False
